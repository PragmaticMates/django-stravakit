"""The activities list as an Excel workbook.

Two sheets: ``Activities`` holds one row per activity in the order the queryset gives
(so the download follows the page's sort), ``Summary`` totals them by year and by sport.
The summary is written as COUNTIF/SUMIF formulas over the first sheet rather than numbers,
so it keeps adding up after the reader deletes or edits rows in Excel.

Dates are the activity's *local* start time — Strava's ``start_date_local``, which is what
the athlete remembers — falling back to the project time zone for rows imported without it.
Speeds arrive from Strava in m/s and leave here in km/h; distance in metres leaves in km.
"""
import io
import unicodedata
from datetime import datetime, timedelta

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

FONT = 'Arial'
HEADER_FONT = Font(name=FONT, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='FC4C02')  # Strava orange
BODY_FONT = Font(name=FONT)
BOLD_FONT = Font(name=FONT, bold=True)
LINK_FONT = Font(name=FONT, color='0563C1', underline='single')
NOTE_FONT = Font(name=FONT, italic=True, color='666666')

# (header, column width, number format). Column letters are referenced by the formulas
# below (C = date, D = year, E = sport, F = km, G = moving time, L = elevation), so a
# column added or moved here has to be followed there.
COLUMNS = (
    ('ID', 12, '0'),
    ('Name', 36, '@'),
    ('Date', 17, 'yyyy-mm-dd hh:mm'),
    ('Year', 7, '0'),
    ('Sport', 18, '@'),
    ('Distance (km)', 14, '0.00'),
    ('Moving time', 13, '[h]:mm:ss'),
    ('Elapsed time', 13, '[h]:mm:ss'),
    ('Pace (min/km)', 14, 'mm:ss'),
    ('Avg speed (km/h)', 16, '0.0'),
    ('Max speed (km/h)', 16, '0.0'),
    ('Elevation (m)', 13, '0'),
    ('Avg HR', 9, '0'),
    ('Max HR', 9, '0'),
    ('Calories', 9, '0'),
    ('Kudos', 8, '0'),
    ('Comments', 10, '0'),
    ('PRs', 6, '0'),
    ('Photos', 8, '0'),
    ('Gear', 28, '@'),
    ('Device', 20, '@'),
    ('Description', 40, '@'),
    ('Strava', 40, '@'),
)
LINK_COLUMN = len(COLUMNS)
ACTIVITIES_SHEET = 'Activities'


def local_start(activity):
    """Naive local start time: Strava's ``start_date_local`` when the stored payload has
    one, else the project time zone's reading of ``start_date``. Naive because Excel has no
    notion of a time zone and openpyxl refuses an aware datetime."""
    raw = (activity.json or {}).get('start_date_local')
    if raw:
        try:
            return datetime.fromisoformat(raw.replace('Z', '+00:00')).replace(tzinfo=None)
        except ValueError:
            pass
    return timezone.localtime(activity.start_date).replace(tzinfo=None)


def _duration(seconds):
    return timedelta(seconds=seconds) if seconds else None


def _kmh(metres_per_second):
    return round(metres_per_second * 3.6, 2) if metres_per_second else None


def _row(activity, row):
    json = activity.json or {}
    return [
        activity.id,
        activity.name,
        local_start(activity),
        f'=YEAR(C{row})',
        str(activity.get_sport_type_display()),
        round(activity.distance / 1000, 3),
        _duration(activity.moving_time),
        _duration(activity.elapsed_time),
        # Duration per kilometre: G is a day fraction, so the quotient renders under mm:ss.
        f'=IF(AND(F{row}>0,G{row}>0),G{row}/F{row},"")',
        _kmh(activity.average_speed),
        _kmh(activity.max_speed),
        activity.total_elevation_gain,
        activity.average_heartrate,
        activity.max_heartrate,
        activity.calories or None,
        activity.kudos_count,
        activity.comment_count,
        activity.pr_count,
        activity.total_photo_count,
        str(activity.gear) if activity.gear else '',
        json.get('device_name') or '',
        json.get('description') or '',
        activity.get_absolute_url(),
    ]


def _write_activities(ws, activities):
    """Fill the activities sheet; returns the number of the last data row."""
    ws.append([header for header, _, _ in COLUMNS])
    for index, (_, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=index)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.row_dimensions[1].height = 30

    row = 1
    for row, activity in enumerate(activities, 2):
        ws.append(_row(activity, row))
        for index, (_, _, number_format) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row, column=index)
            cell.font = BODY_FONT
            cell.number_format = number_format
        link = ws.cell(row=row, column=LINK_COLUMN)
        link.hyperlink = link.value
        link.font = LINK_FONT

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{max(row, 2)}'
    return row


def _write_summary(ws, activities, last_row, note):
    ws.column_dimensions['A'].width = 22
    for column in 'BCDE':
        ws.column_dimensions[column].width = 16

    def column_range(column):
        return f"{ACTIVITIES_SHEET}!${column}$2:${column}${last_row}"

    def header(row, labels):
        for index, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=index, value=label)
            cell.font, cell.fill = HEADER_FONT, HEADER_FILL
            cell.alignment = Alignment(horizontal='left' if index == 1 else 'center')

    def block(start, key_label, keys, key_column):
        """A table of count / km / moving time / elevation per key, plus a total row.
        Returns the row after the gap below it."""
        header(start, [key_label, 'Activities', 'Distance (km)', 'Moving time', 'Elevation (m)'])
        row = start + 1
        for key in keys:
            ws.cell(row=row, column=1, value=key).font = BODY_FONT
            ws.cell(row=row, column=2, value=f'=COUNTIF({column_range(key_column)},A{row})')
            ws.cell(row=row, column=3, value=f'=SUMIF({column_range(key_column)},A{row},{column_range("F")})')
            ws.cell(row=row, column=4, value=f'=SUMIF({column_range(key_column)},A{row},{column_range("G")})')
            ws.cell(row=row, column=5, value=f'=SUMIF({column_range(key_column)},A{row},{column_range("L")})')
            row += 1
        ws.cell(row=row, column=1, value='Total').font = BOLD_FONT
        for column in range(2, 6):
            letter = get_column_letter(column)
            ws.cell(row=row, column=column, value=f'=SUM({letter}{start + 1}:{letter}{row - 1})')
        for current in range(start + 1, row + 1):
            ws.cell(row=current, column=2).number_format = '0'
            ws.cell(row=current, column=3).number_format = '#,##0.0'
            ws.cell(row=current, column=4).number_format = '[h]:mm'
            ws.cell(row=current, column=5).number_format = '#,##0'
            for column in range(2, 6):
                ws.cell(row=current, column=column).font = BOLD_FONT if current == row else BODY_FONT
        return row + 2

    next_row = 1
    if activities:
        years = sorted({local_start(activity).year for activity in activities})
        sports = sorted({str(activity.get_sport_type_display()) for activity in activities})
        next_row = block(next_row, 'Year', years, 'D')
        next_row = block(next_row, 'Sport', sports, 'E')
    else:
        ws.cell(row=next_row, column=1, value='No activities match the filter.').font = BODY_FONT
        next_row += 2
    ws.cell(row=next_row, column=1, value=note).font = NOTE_FONT


def activities_workbook(queryset, athlete=None, note=None):
    """The queryset's activities as .xlsx bytes. ``queryset`` should already be filtered
    and ordered — this writes what it is given, in that order."""
    activities = list(queryset)
    workbook = Workbook()
    activities_sheet = workbook.active
    activities_sheet.title = ACTIVITIES_SHEET
    last_row = _write_activities(activities_sheet, activities)

    if note is None:
        who = f' for {athlete}' if athlete else ''
        note = (f'Source: Strava, exported {timezone.localdate():%Y-%m-%d}{who}. '
                'Dates are local start times; speeds converted from m/s.')
    _write_summary(workbook.create_sheet('Summary'), activities, last_row, note)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_filename(athlete=None):
    """``strava-activities-erik-telepovsky-2026-10-02.xlsx``: the athlete's name folded to
    ASCII, because a Content-Disposition filename is a header and a diacritic in one is
    something every browser handles differently."""
    stem = 'strava-activities'
    if athlete and athlete.full_name:
        folded = unicodedata.normalize('NFKD', athlete.full_name).encode('ascii', 'ignore').decode()
        if folded.strip():
            stem += '-' + '-'.join(folded.lower().split())
    return f'{stem}-{timezone.localdate():%Y-%m-%d}.xlsx'
